"""Reads test data from Excel (Framework step 17: parameterize with Excel)."""
import os

from openpyxl import load_workbook

from utils.config_reader import ROOT


def read_rows(file_name, sheet):
    """Return every data row (skipping the header) as a list of tuples."""
    wb = load_workbook(os.path.join(ROOT, "data", file_name), read_only=True)
    rows = [tuple("" if c is None else str(c) for c in r)
            for r in wb[sheet].iter_rows(min_row=2, values_only=True) if any(r)]
    wb.close()
    return rows
